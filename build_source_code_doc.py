from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "CEDIMS_SourceCode_Copyright_Clean.docx"


SNIPPETS = [
    ("Project Configuration", "package.json", 1, 44),
    ("Authentication and Role Routing", "src/lib/utils/auth.ts", 1, 86),
    ("Document Permission Rules", "src/lib/utils/documentPermissions.ts", 1, 78),
    ("Upload Page Main State", "src/routes/dashboard/upload/+page.svelte", 1, 80),
    ("Upload Page Processing", "src/routes/dashboard/upload/+page.svelte", 125, 230),
    ("Upload Pipeline Core", "src/lib/utils/pipeline.ts", 1, 108),
    ("Online and Offline Pipeline", "src/lib/utils/pipeline.ts", 170, 290),
    ("OCR Metadata Extraction", "src/lib/utils/ocr.ts", 1, 115),
    ("OCR Parsing", "src/lib/utils/ocr.ts", 175, 285),
    ("Fuzzy Classifier", "src/lib/utils/fuzzyClassifier.ts", 1, 92),
    ("Offline Cache and Queue", "src/lib/utils/offline.ts", 1, 112),
    ("Offline Sync Processor", "src/lib/utils/offline.ts", 300, 420),
    ("Analytics Queries", "src/lib/utils/analyticsQueries.ts", 1, 120),
    ("Storage Upload API", "src/routes/api/storage/upload/+server.ts", 1, 70),
    ("Storage Presign API", "src/routes/api/storage/presign/+server.ts", 1, 95),
    ("Verification Page", "src/routes/verify/[hash]/+page.svelte", 1, 135),
    ("Dashboard Main Page", "src/routes/dashboard/+page.svelte", 1, 150),
    ("Compliance Dashboard", "src/routes/dashboard/compliance/+page.svelte", 1, 150),
    ("Archive Page", "src/routes/dashboard/archive/+page.svelte", 1, 130),
]


def set_run_font(run, name="Arial", size=9, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)


def set_columns(section, count=2, space="360"):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols = cols[0]
    else:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), space)


def set_document_defaults(doc):
    styles = doc.styles
    for style_name in ["Normal", "Body Text"]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(9)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

    for style_name, size in [("Title", 14), ("Heading 1", 11), ("Heading 2", 10)]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")


def read_lines(rel, start, end):
    path = ROOT / rel
    if not path.exists():
        return [f"// File not found: {rel}"]
    text = path.read_text(encoding="utf-8", errors="replace").splitlines()
    lines = text[start - 1 : min(end, len(text))]
    return strip_comments(lines)


def strip_comments(lines):
    cleaned = []
    in_block = False
    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()

        if in_block:
            if "*/" in stripped:
                in_block = False
                remainder = stripped.split("*/", 1)[1].strip()
                if remainder and not remainder.startswith("//"):
                    cleaned.append(remainder)
            continue

        if not stripped:
            cleaned.append(line)
            continue

        if stripped.startswith("//") or stripped.startswith("--"):
            continue

        if stripped.startswith("/*"):
            if "*/" not in stripped:
                in_block = True
                continue
            after = stripped.split("*/", 1)[1].strip()
            if after:
                cleaned.append(after)
            continue

        if stripped.startswith("*"):
            continue

        cleaned.append(strip_inline_comment(line))

    return cleaned


def strip_inline_comment(line):
    in_single = False
    in_double = False
    in_backtick = False
    escaped = False
    for i in range(len(line) - 1):
        ch = line[i]
        nxt = line[i + 1]
        if escaped:
            escaped = False
            continue
        if ch == "\\":
            escaped = True
            continue
        if ch == "'" and not in_double and not in_backtick:
            in_single = not in_single
            continue
        if ch == '"' and not in_single and not in_backtick:
            in_double = not in_double
            continue
        if ch == "`" and not in_single and not in_double:
            in_backtick = not in_backtick
            continue
        if ch == "/" and nxt == "/" and not in_single and not in_double and not in_backtick:
            return line[:i].rstrip()
    return line


def add_para(doc, text="", style=None, size=9, bold=False, before=0, after=0):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.0
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line[:118])
        set_run_font(run, size=9)


def build():
    doc = Document()
    set_document_defaults(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("CEDIMS_SourceCode")
    set_run_font(run, size=14, bold=True)

    add_para(
        doc,
        "Coding languages used: Svelte, TypeScript, JavaScript, SQL, CSS, and Python. "
        "Frameworks and services include SvelteKit, Supabase, Backblaze B2 compatible S3 storage, "
        "Tesseract.js, PDF.js, and Vite.",
        size=9,
        after=4,
    )
    add_para(
        doc,
        "This source-code document includes only the main copyright-relevant project features: authentication, "
        "role-based permissions, document upload and OCR, duplicate protection, offline sync, analytics, verification, "
        "storage APIs, dashboards, and archive management.",
        size=9,
        after=8,
    )

    doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(doc.sections[-1], 2)

    for title_text, rel, start, end in SNIPPETS:
        heading = doc.add_paragraph(style="Heading 1")
        heading.paragraph_format.space_before = Pt(5)
        heading.paragraph_format.space_after = Pt(2)
        r = heading.add_run(title_text)
        set_run_font(r, size=10, bold=True)

        meta = doc.add_paragraph()
        meta.paragraph_format.space_after = Pt(2)
        mr = meta.add_run(f"{rel}  Lines {start}-{end}")
        set_run_font(mr, size=9, bold=True)

        add_code_block(doc, read_lines(rel, start, end))

    doc.save(OUT)


if __name__ == "__main__":
    build()
