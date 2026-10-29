from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r'C:\Users\ASUS\OneDrive\Documents\ChatGPT\CDEDIMSCALAPAN\SMART-E_Vision_Calapan_East_District_Instructional_Monitoring_System.docx'
doc = Document(); sec = doc.sections[0]
sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=doc.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5)
for n,s in [('Title',24),('Heading 1',16),('Heading 2',12.5),('Heading 3',11)]:
    st=styles[n]; st.font.name='Aptos Display'; st.font.size=Pt(s); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)

def shade(cell, fill):
    tc=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tc.append(shd)
def borders(t):
    p=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for e in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); b.append(x)
    p.append(b)
def add_table(rows):
    t=doc.add_table(rows=0,cols=len(rows[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for q in cells[i].paragraphs: q.paragraph_format.space_before=Pt(2); q.paragraph_format.space_after=Pt(2)
            if ri==0:
                shade(cells[i],'1F4E79')
                for r in cells[i].paragraphs[0].runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
            elif ri%2==0: shade(cells[i],'F3F6F9')
    return t
def para(text='', align=None, bold=False):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6)
    if align is not None: p.alignment=align
    p.add_run(text).bold=bold; return p
def bullets(items):
    for x in items: doc.add_paragraph(x, style='List Bullet')
def h(text, level=1): doc.add_heading(text, level=level)

doc.add_paragraph().add_run().add_break()
t=doc.add_paragraph(style='Title'); t.alignment=WD_ALIGN_PARAGRAPH.CENTER; t.add_run('SMART-E Vision')
para('Calapan East District Instructional Monitoring System', WD_ALIGN_PARAGRAPH.CENTER, True)
para('Capstone System Documentation', WD_ALIGN_PARAGRAPH.CENTER)
para('Official Capstone Title: SMART-E Vision\nFormal System Name: Calapan East District Instructional Monitoring System\nPrepared for Calapan East District\n2026', WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

h('Document Purpose'); para('This document presents the functional scope, artificial intelligence features, machine-learning methods, analytics capabilities, technology stack, architecture, security controls, and deployment environment of SMART-E Vision. The system is designed for Calapan East District and is deployed as the Calapan East District Instructional Monitoring System.')
h('System Identity'); add_table([['Item','Official description'],['Capstone title','SMART-E Vision'],['Formal system name','Calapan East District Instructional Monitoring System'],['Primary users','Teachers, Master Teachers, School Heads, District Supervisors, and Administrators'],['Purpose','Digital submission, monitoring, review, archiving, analytics, and AI-assisted instructional support'],['Production hosting','OVH Ubuntu VPS with Docker and Docker Compose'],['Database','Supabase PostgreSQL, Authentication, and Realtime'],['Object storage','Backblaze B2 through an S3-compatible interface']])
h('1. System Overview'); para('SMART-E Vision is an AI-assisted web-based instructional monitoring system developed for Calapan East District. It digitizes the submission, verification, review, storage, and analysis of instructional documents. The platform connects teacher submissions with teaching loads, academic-calendar deadlines, review records, and district monitoring requirements.')
para('The system combines deterministic compliance rules with artificial intelligence and machine-learning components. OCR and fuzzy matching assist with document metadata, a supervised intent classifier powers the Gabay assistant, and K-Means clustering groups teacher or school performance patterns for supervisory decision support.')
h('2. General Objective'); para('To develop an intelligent and integrated instructional monitoring system that improves the submission, verification, review, archiving, and analysis of instructional documents in Calapan East District.')
h('3. Specific Objectives'); bullets(['Digitize the submission and storage of instructional documents.','Monitor teacher and school compliance with instructional-document requirements.','Identify missing, late, upcoming, submitted, and reviewed documents.','Reduce manual encoding and repetitive monitoring work.','Apply OCR and fuzzy matching to assist with document metadata extraction.','Use supervised machine learning to classify user questions for the Gabay assistant.','Use K-Means clustering to identify instructional compliance patterns.','Provide role-specific dashboards and reports for school and district supervision.','Support QR-based verification and SHA-256 document integrity checking.','Support offline document preparation and later synchronization.'])

h('4. Target Users and Responsibilities'); add_table([['User role','Main capabilities'],['Teacher','Upload documents, review personal compliance, confirm metadata suggestions, view remarks, and synchronize pending offline submissions.'],['Master Teacher','Review assigned submissions, add remarks, monitor teachers, and follow up missing or late documents.'],['School Head','Monitor school submissions, review documents, compare teacher performance, and export reports.'],['District Supervisor','Monitor schools, compare performance, analyze district compliance, review clusters, and prioritize support.'],['Administrator','Manage accounts, schools, teaching loads, calendars, settings, audit records, and permissions.']])
h('5. Core Functional Modules')
h('5.1 Authentication and Role-Based Access',2); para('The system uses authenticated accounts and role-based permissions. Access is scoped by school and district assignment so users see records appropriate to their responsibilities. Reviewer, teacher, school, district, and administrator actions are controlled by application permissions and database policies.')
h('5.2 Instructional Document Submission',2); para('The upload module supports DLL/DLP, ISP, ISR, and supplementary instructional documents. A submission may pass through validation, conversion, OCR, metadata prediction, teaching-load matching, deadline matching, duplicate detection, hashing, QR stamping, cloud upload, and database registration.')
h('5.3 Compliance Monitoring',2); para('Compliance is calculated by comparing actual submissions with expected requirements from teaching-load and academic-calendar records. The system distinguishes on-time, late, missing, upcoming, approved-leave, for-checking, checked, and supplementary records.')
add_table([['Metric','Meaning'],['Expected','Requirements defined by active teaching loads and academic-calendar periods.'],['Submitted','Requirements covered by an eligible instructional document.'],['Missing','Due requirements without a qualifying submission.'],['On-time','Submitted before the applicable deadline.'],['Late','Submitted after the applicable deadline.'],['For Checking','Reviewable document without a valid reviewer remark.'],['Checked','Reviewable document with a valid reviewer remark.'],['Completion rate','Submitted requirements divided by expected requirements.']])
h('5.4 Digital Archive and Review Tracker',2); para('The archive provides folder navigation, search, document sorting, preview, download, review filtering, reviewer remarks, and permission-controlled management. The Review Tracker provides a focused queue of reviewable documents and maintains the same review-status rules used by Compliance and Archives.')
h('5.5 Academic Calendar and Teaching Loads',2); para('Academic-calendar entries define terms, weeks, deadlines, and active periods. Teaching-load entries associate teachers with subjects and grade levels. Together they determine expected instructional documents and submission evaluation.')
h('5.6 QR Verification and Document Integrity',2); para('Submitted files may contain QR verification data. SHA-256 hashes support file identity and duplicate checking. A verification route can display document metadata and verification results. Cached verification data supports limited offline checking.')
h('5.7 Offline-First Operation',2); para('The system uses browser storage, a service worker, and a synchronization queue for unstable connections. Documents can be prepared while offline, placed in a pending queue, and synchronized when connectivity returns. External database requests are excluded from the service-worker offline fallback.')

doc.add_page_break(); h('6. Artificial Intelligence Features')
h('6.1 OCR-Assisted Metadata Extraction',2); para('Tesseract.js and PDF.js extract text from uploaded instructional documents. The extracted text assists with subject, grade level, document type, term, week, and instructional-heading detection. OCR confidence is exposed so low-confidence results can be verified by the user.')
h('6.2 Fuzzy Metadata Classifier',2); para('OCR results may contain spelling errors, broken words, inconsistent capitalization, or scanning noise. The fuzzy classifier compares extracted text with known subjects, grade levels, and document types using character-level similarity and Dice/bigram comparison. It returns a likely value and confidence score and can compare the result with teaching loads.')
h('6.3 Gabay AI Assistant',2); para('Gabay is a domain-specific AI assistant for SMART-E Vision. It answers questions about uploads, compliance, deadlines, archives, QR verification, reports, teaching loads, school performance, district performance, and system features. It combines intent classification, fuzzy keyword matching, slot extraction, confidence scoring, knowledge-base responses, and live application data where appropriate.')
para('Gabay is documented as a hybrid, domain-specific AI assistant. It is not described as a general-purpose large language model because the current implementation uses a trained local intent model and structured system knowledge responses.')
h('7. Machine Learning Features')
h('7.1 Supervised Intent Classification',2); para('The chatbot intent model is trained using Python, pandas, and scikit-learn. Character n-gram features are converted into numerical representations and classified using Logistic Regression. The training workflow includes train/test splitting, stratified cross-validation, classification reports, confusion matrices, ROC analysis, and confidence scoring.')
add_table([['Component','Implementation'],['Learning type','Supervised classification'],['Model','Logistic Regression'],['Features','Character n-grams'],['Training tools','Python, pandas, and scikit-learn'],['Output','JSON model with vocabulary, coefficients, labels, and metrics'],['Use','Routing user questions to the correct system response']])
h('7.2 K-Means Compliance Clustering',2); para('K-Means is applied to teacher or school feature vectors derived from compliance and submission behavior. Features may include completion, missing documents, late and on-time submissions, review completion, and weekly patterns. Groups are labeled for interpretation, such as Consistently Meeting Standards, Steadily Progressing, and Building Momentum.')
para('Clustering is a decision-support feature. It helps supervisors identify patterns and prioritize follow-up; it does not independently make disciplinary or administrative decisions.')

h('8. Analytics and Reporting'); bullets(['District-wide compliance and submission summaries.','School-level comparison and performance analysis.','Teacher-level submission and review statistics.','Missing and late submission monitoring.','Review completion and reviewer-workload monitoring.','Compliance trend and performance-distribution views.','K-Means cluster summaries and membership filters.','Risk and priority indicators for supervisory follow-up.','Excel and Word report generation for compliance and review records.'])
h('9. Technology Stack'); add_table([['Layer','Technology'],['Frontend','Svelte 5, SvelteKit, TypeScript, Vite, Tailwind CSS, and Lucide Svelte'],['Backend runtime','Node.js 22, SvelteKit adapter-node, and SvelteKit server routes'],['Database','Supabase PostgreSQL'],['Authentication','Supabase Auth'],['Realtime','Supabase Realtime subscriptions'],['Cloud storage','Backblaze B2 through an S3-compatible API'],['OCR and PDF','Tesseract.js, PDF.js, pdf-lib, pdf-parse, and Mammoth'],['AI and ML','scikit-learn Logistic Regression, character n-grams, fuzzy Dice/bigram matching, and K-Means'],['Reports','ExcelJS and docx'],['Verification','QRCode and jsQR with SHA-256 hashing'],['Deployment','OVH Ubuntu 24.04 VPS, Docker, and Docker Compose'],['Offline support','Service Worker, IndexedDB, and background synchronization']])

doc.add_page_break(); h('10. System Architecture')
para('The system follows a cloud-connected, containerized web architecture. The browser provides the user interface and performs selected OCR, fuzzy matching, offline queueing, and visualization tasks. SvelteKit server routes handle protected operations and storage access. Supabase provides authentication, PostgreSQL data, and realtime events. Backblaze B2 stores uploaded documents. The complete production application runs in a Docker container on an OVH VPS.')
add_table([['Component','Responsibility'],['Browser client','User interface, OCR assistance, fuzzy matching, offline queue, dashboards, charts, and QR scanning.'],['SvelteKit application','Protected server routes, storage presigning, document operations, report support, and application delivery.'],['Supabase','Authentication, PostgreSQL records, row-level security, realtime updates, and audit data.'],['Backblaze B2','Durable storage for instructional documents.'],['OVH VPS','Production runtime for the Dockerized SvelteKit application.']])
h('11. Main Data Flow'); para('A teacher selects an instructional document. The system validates the file, extracts text through OCR, predicts metadata, applies fuzzy matching, checks the teaching load and academic period, calculates submission status, computes the SHA-256 identity, adds verification data, stores the file in Backblaze B2, and records the submission in Supabase. The record becomes available to Archive, Compliance, Review Tracker, Analytics, and reporting modules.')
h('12. Security and Data Protection'); bullets(['Authenticated sessions and role-based authorization.','School and district data scoping.','Supabase Row-Level Security policies.','Protected document access through presigned URLs.','SHA-256 file hashing and QR-based verification.','Audit logging for important activities.','Controlled reviewer remarks and administrative actions.','Production secrets stored in VPS environment variables.','HTTPS deployment and restricted container binding.','No direct exposure of storage credentials to the browser.'])
h('13. Deployment and Hosting'); para('Production uses an OVH Ubuntu 24.04 VPS. The application is built and served through a Node.js 22 Docker image managed by Docker Compose. The container listens internally on port 3000 and is exposed locally to the reverse proxy. Docker restarts the service unless stopped and performs a health check against the application root route.')
add_table([['Deployment element','Current setup'],['Server','OVH VPS'],['Operating system','Ubuntu 24.04'],['Runtime','Node.js 22'],['Containerization','Docker and Docker Compose'],['Application port','127.0.0.1:3000 inside the VPS'],['Database','Supabase cloud project'],['Storage','Backblaze B2'],['Source delivery','GitHub main branch pulled into the application deployment directory'],['Operational check','Docker health check and HTTP verification']])
h('14. Functional Features Confirmed'); bullets(['User authentication and role-based access.','Teacher, school, district, and administrator workflows.','Teaching-load and academic-calendar management.','DLL/DLP, ISP, ISR, and supplementary document submissions.','OCR-assisted metadata extraction.','Fuzzy subject, grade, and document-type matching.','Deadline, missing, late, and on-time classification.','Digital archives with search, folders, sorting, and filters.','For Checking and Checked review workflow.','Reviewer remarks and reviewer identity.','QR creation, QR scanning, and public verification.','SHA-256 document hashing and duplicate detection.','Offline submission queue and background synchronization.','Supabase Realtime updates.','Gabay AI assistant and intent classification.','K-Means performance clustering.','Compliance dashboards and analytics.','Excel and Word report generation.','Audit logs, notifications, responsive UI, and dark mode.'])
h('15. Documentation Accuracy Notes'); para('SMART-E Vision should be used as the capstone title throughout the research manuscript, proposal, presentation, and defense documentation. Calapan East District Instructional Monitoring System should be used as the formal system name throughout the documentation and deployment interface.')
para('The system combines artificial intelligence, machine learning, deterministic compliance rules, and database analytics. Compliance rates are calculated from actual submission and academic-calendar records. OCR and fuzzy predictions assist users but may require confirmation. K-Means clustering summarizes patterns and supports supervisory decisions; it does not make independent administrative decisions.')
h('16. Recommended Capstone Description'); para('SMART-E Vision: Calapan East District Instructional Monitoring System is an AI-assisted web-based platform developed to improve the submission, verification, monitoring, review, archiving, and analysis of instructional documents within Calapan East District. The system integrates OCR, fuzzy metadata classification, supervised intent classification, K-Means clustering, QR-based verification, SHA-256 hashing, offline-first synchronization, role-based access control, realtime database updates, cloud storage, and analytics dashboards for evidence-based instructional supervision.')

for s in doc.sections:
    f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; f.add_run('SMART-E Vision | Calapan East District Instructional Monitoring System').font.size=Pt(8)
doc.core_properties.title='SMART-E Vision Calapan East District Instructional Monitoring System'; doc.core_properties.subject='Capstone system documentation'; doc.core_properties.author='Calapan East District'
doc.save(OUT); print(OUT)
